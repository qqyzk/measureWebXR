from __future__ import annotations

import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd


WORKBOOK_PATH = Path(r"D:\TMC\result.xlsx")
OUTPUT_DIR = Path(r"D:\TMC\measureWebXR\results\scene_plots")

FRAMEWORK_ORDER = ["three", "baby", "Aframe", "play", "react"]
FRAMEWORK_LABELS = {
    "three": "Three.js",
    "baby": "Babylon.js",
    "Aframe": "A-Frame",
    "play": "PlayCanvas",
    "react": "React Three Fiber",
}
FRAMEWORK_COLORS = {
    "three": "#4C78A8",
    "baby": "#F58518",
    "Aframe": "#54A24B",
    "play": "#E45756",
    "react": "#8F63C7",
}
DEPTH_COLORS = {
    1: "#4C78A8",
    2: "#F58518",
    3: "#54A24B",
    4: "#E45756",
    5: "#8F63C7",
}
METRICS = {
    "fps": {
        "title": "Scene 4096 FPS by Depth",
        "ylabel": "FPS",
        "yscale": "linear",
        "filename": "scene_fps.png",
    },
    "load": {
        "title": "Scene 4096 Load Time by Depth",
        "ylabel": "Load Time (ms)",
        "yscale": "log",
        "filename": "scene_load_time.png",
    },
    "ft": {
        "title": "Scene 4096 FLT by Depth",
        "ylabel": "FLT (ms)",
        "yscale": "log",
        "filename": "scene_ft.png",
    },
}
SECTION_NAMES = ("multi", "instance")
CHANGE_METRICS = {
    "fps": {
        "title": "Instance vs Multi FPS Ratio by Framework",
        "ylabel": "Instance / Multi (%)",
        "filename": "instance_vs_multi_fps_ratio.png",
    },
    "ft": {
        "title": "Instance vs Multi FLT Ratio by Framework",
        "ylabel": "Instance / Multi (%)",
        "filename": "instance_vs_multi_ft_ratio.png",
    },
}
COUNT_COLORS = {
    2: "#4C78A8",
    4: "#F58518",
    8: "#54A24B",
    16: "#E45756",
    32: "#8F63C7",
}
LIGHTING_RATIO_METRICS = {
    "fps": {
        "title": "Phong vs PBR FPS Ratio by Framework",
        "ylabel": "Phong / PBR (%)",
        "filename": "lighting_phong_vs_pbr_fps_ratio.png",
    },
    "ft": {
        "title": "Phong vs PBR FLT Ratio by Framework",
        "ylabel": "Phong / PBR (%)",
        "filename": "lighting_phong_vs_pbr_ft_ratio.png",
    },
}


def to_number(value):
    if value is None or isinstance(value, bool):
        return math.nan
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return math.nan


def extract_scene_rows(workbook_path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    records = []

    for sheet_name in FRAMEWORK_ORDER:
        ws = wb[sheet_name]
        scene_row = None
        for row_idx, (label, *_rest) in enumerate(ws.iter_rows(values_only=True), start=1):
            if isinstance(label, str) and label.strip().lower().startswith("scene "):
                scene_row = row_idx
                break

        if scene_row is None:
            raise ValueError(f"Scene section not found in sheet: {sheet_name}")

        row_idx = scene_row + 1
        while row_idx <= ws.max_row:
            label = ws.cell(row=row_idx, column=1).value
            if not isinstance(label, str) or not label.strip().lower().startswith("d="):
                break

            depth = int(label.split("=", 1)[1])
            records.append(
                {
                    "framework": sheet_name,
                    "framework_label": FRAMEWORK_LABELS[sheet_name],
                    "depth": depth,
                    "load": to_number(ws.cell(row=row_idx, column=2).value),
                    "fps": to_number(ws.cell(row=row_idx, column=3).value),
                    "ft": to_number(ws.cell(row=row_idx, column=4).value),
                }
            )
            row_idx += 1

    return pd.DataFrame.from_records(records)


def extract_named_block_rows(workbook_path: Path, block_name: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    records = []

    for sheet_name in FRAMEWORK_ORDER:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            label = row[0]
            if not isinstance(label, str):
                continue

            parts = label.strip().split()
            if len(parts) != 2 or parts[0].lower() != block_name:
                continue

            count = int(parts[1])
            records.append(
                {
                    "framework": sheet_name,
                    "framework_label": FRAMEWORK_LABELS[sheet_name],
                    "count": count,
                    "load": to_number(row[1]),
                    "fps": to_number(row[2]),
                    "ft": to_number(row[3]),
                }
            )

    return pd.DataFrame.from_records(records)


def extract_lighting_rows(workbook_path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    records = []

    for sheet_name in FRAMEWORK_ORDER:
        ws = wb[sheet_name]
        lighting_row = None
        for row_idx, (label, *_rest) in enumerate(ws.iter_rows(values_only=True), start=1):
            if isinstance(label, str) and label.strip().lower() == "lighting":
                lighting_row = row_idx
                break

        if lighting_row is None:
            continue

        row_idx = lighting_row + 1
        while row_idx <= ws.max_row:
            label = ws.cell(row=row_idx, column=1).value
            if isinstance(label, str):
                row_idx += 1
                continue
            if label is None:
                break

            count = int(label)
            records.append(
                {
                    "framework": sheet_name,
                    "framework_label": FRAMEWORK_LABELS[sheet_name],
                    "count": count,
                    "load": to_number(ws.cell(row=row_idx, column=2).value),
                    "fps": to_number(ws.cell(row=row_idx, column=3).value),
                    "ft": to_number(ws.cell(row=row_idx, column=4).value),
                }
            )
            row_idx += 1

    return pd.DataFrame.from_records(records)


def calculate_change_rates(multi_df: pd.DataFrame, instance_df: pd.DataFrame) -> pd.DataFrame:
    merged = multi_df.merge(
        instance_df,
        on=["framework", "framework_label", "count"],
        suffixes=("_multi", "_instance"),
    )

    for metric in ("fps", "ft"):
        baseline = merged[f"{metric}_multi"]
        compared = merged[f"{metric}_instance"]
        merged[f"{metric}_ratio_pct"] = np.where(
            baseline.notna() & compared.notna() & (baseline != 0),
            compared / baseline * 100.0,
            np.nan,
        )

    return merged.sort_values(["count", "framework"]).reset_index(drop=True)


def calculate_ratio_rates(
    baseline_df: pd.DataFrame,
    compared_df: pd.DataFrame,
    baseline_suffix: str,
    compared_suffix: str,
) -> pd.DataFrame:
    merged = baseline_df.merge(
        compared_df,
        on=["framework", "framework_label", "count"],
        suffixes=(f"_{baseline_suffix}", f"_{compared_suffix}"),
    )

    for metric in ("fps", "ft"):
        baseline = merged[f"{metric}_{baseline_suffix}"]
        compared = merged[f"{metric}_{compared_suffix}"]
        merged[f"{metric}_ratio_pct"] = np.where(
            baseline.notna() & compared.notna() & (baseline != 0),
            compared / baseline * 100.0,
            np.nan,
        )

    return merged.sort_values(["count", "framework"]).reset_index(drop=True)


def plot_metric(df: pd.DataFrame, metric: str, config: dict[str, str], output_dir: Path) -> None:
    depths = sorted(df["depth"].unique())
    x = np.arange(len(FRAMEWORK_ORDER), dtype=float)
    total_width = 0.82
    bar_width = total_width / len(depths)

    fig, ax = plt.subplots(figsize=(10.5, 6.5))

    for idx, depth in enumerate(depths):
        framework_df = (
            df[df["depth"] == depth]
            .set_index("framework")
            .reindex(FRAMEWORK_ORDER)
        )
        values = framework_df[metric].to_numpy(dtype=float)
        positions = x - total_width / 2 + (idx + 0.5) * bar_width
        ax.bar(
            positions,
            values,
            width=bar_width,
            label=f"d={depth}",
            color=DEPTH_COLORS[depth],
            edgecolor="black",
            linewidth=0.6,
        )

    ax.set_title(config["title"], fontsize=20, pad=14)
    ax.set_xlabel("Framework", fontsize=15)
    ax.set_ylabel(config["ylabel"], fontsize=15)
    ax.set_xticks(x)
    ax.set_xticklabels([FRAMEWORK_LABELS[framework] for framework in FRAMEWORK_ORDER], fontsize=12)
    ax.tick_params(axis="y", labelsize=12)
    ax.set_yscale(config["yscale"])
    ax.grid(False, axis="x")
    ax.grid(axis="y", linestyle="--", linewidth=0.7, alpha=0.45)
    ax.set_axisbelow(True)
    ax.legend(title="Scene Depth", ncol=3, frameon=False, fontsize=11, title_fontsize=11)

    fig.tight_layout()
    fig.savefig(output_dir / config["filename"], dpi=220, bbox_inches="tight")
    plt.close(fig)


def plot_change_metric(df: pd.DataFrame, metric: str, config: dict[str, str], output_dir: Path) -> None:
    counts = sorted(df["count"].unique())
    x = np.arange(len(FRAMEWORK_ORDER), dtype=float)
    total_width = 0.82
    bar_width = total_width / len(counts)

    fig, ax = plt.subplots(figsize=(10.5, 6.5))

    for idx, count in enumerate(counts):
        framework_df = (
            df[df["count"] == count]
            .set_index("framework")
            .reindex(FRAMEWORK_ORDER)
        )
        values = framework_df[f"{metric}_ratio_pct"].to_numpy(dtype=float)
        positions = x - total_width / 2 + (idx + 0.5) * bar_width
        ax.bar(
            positions,
            values,
            width=bar_width,
            label=f"n={count}",
            color=COUNT_COLORS[count],
            edgecolor="black",
            linewidth=0.6,
        )

    ax.set_title(config["title"], fontsize=20, pad=14)
    ax.set_xlabel("Framework", fontsize=15)
    ax.set_ylabel(config["ylabel"], fontsize=15)
    ax.set_xticks(x)
    ax.set_xticklabels([FRAMEWORK_LABELS[framework] for framework in FRAMEWORK_ORDER], fontsize=12)
    ax.tick_params(axis="y", labelsize=12)
    ax.grid(False, axis="x")
    ax.axhline(100, color="black", linewidth=1)
    ax.grid(axis="y", linestyle="--", linewidth=0.7, alpha=0.45)
    ax.set_axisbelow(True)
    ax.legend(title="Model Count", ncol=3, frameon=False, fontsize=11, title_fontsize=11)

    fig.tight_layout()
    fig.savefig(output_dir / config["filename"], dpi=220, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    df = extract_scene_rows(WORKBOOK_PATH)
    df.sort_values(["depth", "framework"], inplace=True)
    df.to_csv(OUTPUT_DIR / "scene_metrics.csv", index=False)

    plt.style.use("seaborn-v0_8-whitegrid")
    for metric, config in METRICS.items():
        plot_metric(df, metric, config, OUTPUT_DIR)

    multi_df = extract_named_block_rows(WORKBOOK_PATH, "multi")
    instance_df = extract_named_block_rows(WORKBOOK_PATH, "instance")
    lighting_df = extract_lighting_rows(WORKBOOK_PATH)
    multi_df.sort_values(["count", "framework"], inplace=True)
    instance_df.sort_values(["count", "framework"], inplace=True)
    lighting_df.sort_values(["count", "framework"], inplace=True)
    multi_df.to_csv(OUTPUT_DIR / "multi_metrics.csv", index=False)
    instance_df.to_csv(OUTPUT_DIR / "instance_metrics.csv", index=False)
    lighting_df.to_csv(OUTPUT_DIR / "lighting_metrics.csv", index=False)

    change_df = calculate_ratio_rates(multi_df, instance_df, "multi", "instance")
    change_df.to_csv(OUTPUT_DIR / "instance_vs_multi_change.csv", index=False)
    for metric, config in CHANGE_METRICS.items():
        plot_change_metric(change_df, metric, config, OUTPUT_DIR)

    lighting_ratio_df = calculate_ratio_rates(multi_df, lighting_df, "pbr", "phong")
    lighting_ratio_df.to_csv(OUTPUT_DIR / "lighting_phong_vs_pbr_ratio.csv", index=False)
    for metric, config in LIGHTING_RATIO_METRICS.items():
        plot_change_metric(lighting_ratio_df, metric, config, OUTPUT_DIR)

    print(f"Saved plots to: {OUTPUT_DIR}")
    print("Scene metrics")
    print(df.to_string(index=False))
    print("Instance vs multi changes")
    print(
        change_df[
            [
                "framework",
                "framework_label",
                "count",
                "fps_multi",
                "fps_instance",
                "fps_ratio_pct",
                "ft_multi",
                "ft_instance",
                "ft_ratio_pct",
            ]
        ].to_string(index=False)
    )
    print("Lighting phong vs pbr ratios")
    print(
        lighting_ratio_df[
            [
                "framework",
                "framework_label",
                "count",
                "fps_pbr",
                "fps_phong",
                "fps_ratio_pct",
                "ft_pbr",
                "ft_phong",
                "ft_ratio_pct",
            ]
        ].to_string(index=False)
    )


if __name__ == "__main__":
    main()
